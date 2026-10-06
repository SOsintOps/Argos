# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Ramingo (SOsintOps)
"""Build the Argos report templates (.docx) in the templates/ folder.

Usage (needs python-docx):
    uv run --with python-docx templates/src/build_templates.py
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent
GREY = RGBColor(0x59, 0x59, 0x59)
ACCENT = RGBColor(0x0F, 0x76, 0x6E)

CONFIDENCE = (
    "Confidence: High = several independent, reliable sources agree; "
    "Moderate = credible but not fully corroborated; Low = single or "
    "questionable source, treat as a lead. Likelihood words: almost no chance, "
    "very unlikely, unlikely, roughly even chance, likely, very likely, almost certain."
)
ADMIRALTY = (
    "Source reliability A (reliable) to F (cannot be judged); information "
    "credibility 1 (confirmed by other sources) to 6 (cannot be judged)."
)


def new_document(title, subtitle):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Liberation Sans"
    style.font.size = Pt(10.5)
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "CLASSIFICATION / HANDLING: ________________"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer = section.footer.paragraphs[0]
        footer.text = "Prepared with Argos — keep with the case evidence"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph(subtitle)
    p.runs[0].font.color.rgb = ACCENT
    return doc


def note(doc, text):
    """Guidance for the author: grey italic, to delete before release."""
    p = doc.add_paragraph()
    run = p.add_run("[" + text + "]")
    run.italic = True
    run.font.color.rgb = GREY
    return p


def shade(cell, color="E7F2F1"):
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:color"), "auto")
    fill.set(qn("w:fill"), color)
    props.append(fill)


def fields(doc, labels, cols=2):
    """Two-column (or four-column) table of labels with empty value cells."""
    per_row = cols // 2
    rows = (len(labels) + per_row - 1) // per_row
    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    for i, label in enumerate(labels):
        r, c = divmod(i, per_row)
        cell = table.cell(r, c * 2)
        cell.text = label
        cell.paragraphs[0].runs[0].bold = True
        shade(cell)
    doc.add_paragraph()
    return table


def grid(doc, headers, rows=4):
    table = doc.add_table(rows=rows + 1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        shade(cell)
    doc.add_paragraph()
    return table


def checkboxes(doc, items):
    for item in items:
        doc.add_paragraph("☐  " + item)


def sources_annex(doc):
    doc.add_heading("Sources and evidence", level=1)
    note(doc, "One row per source. Run folder and SHA-256 come from the Argos run "
              "folder (command.txt and SHA256SUMS), so every statement can be traced "
              "back to the exact capture. " + ADMIRALTY)
    grid(doc, ["Ref", "Description", "URL / origin", "Captured (UTC)",
               "Argos run folder", "SHA-256", "Rel./Cred."], rows=6)


def save(doc, name):
    doc.save(OUT / name)
    print("written", name)


# ── Templates ────────────────────────────────────────────────────────────────

def full_report():
    doc = new_document("Investigation Report", "Open-source intelligence — full report")
    fields(doc, ["Case", "Report reference", "Subject(s)", "Requested by", "Prepared by",
                 "Reviewed by", "Date of report", "Period covered"], cols=4)
    doc.add_heading("1. Bottom line", level=1)
    note(doc, "Three to five sentences a busy reader can act on: what was asked, "
              "what was found, how confident you are, what you recommend.")
    doc.add_heading("2. Tasking and scope", level=1)
    note(doc, "The question you were asked, the legal basis or authorisation, what "
              "was in scope and out of scope, and any restrictions (no contact, "
              "no account creation, passive only).")
    doc.add_heading("3. Key findings", level=1)
    note(doc, CONFIDENCE)
    grid(doc, ["#", "Finding", "Confidence", "Sources (Ref)"], rows=5)
    doc.add_heading("4. Subject overview", level=1)
    fields(doc, ["Name(s) and aliases", "Date / place of birth", "Nationality",
                 "Addresses", "Phone numbers", "Email addresses", "Usernames",
                 "Employer / role", "Vehicles", "Associates"])
    doc.add_heading("Online presence", level=2)
    grid(doc, ["Platform", "Handle / URL", "Account ID", "Status", "Ref"], rows=6)
    doc.add_heading("5. Narrative", level=1)
    note(doc, "Chronological account of the research: what you looked for, where, "
              "and what each step showed. Mention dead ends too: they show coverage.")
    doc.add_heading("6. Timeline", level=1)
    grid(doc, ["Date / time (UTC)", "Event", "Ref"], rows=6)
    doc.add_heading("7. Analysis", level=1)
    note(doc, "Links between people, accounts and infrastructure; alternative "
              "explanations you considered and why you discarded them; gaps that "
              "remain and how they could be closed.")
    doc.add_heading("8. Recommendations", level=1)
    note(doc, "Next steps, further collection, referrals.")
    sources_annex(doc)
    doc.add_heading("Method and limitations", level=1)
    note(doc, "Workstation and tools used (Argos version), research identity and "
              "network used, date range of collection, and the limits of open "
              "sources for this question.")
    save(doc, "Argos_Report_Full.docx")


def executive_report():
    doc = new_document("Executive Summary", "Open-source intelligence — short report")
    fields(doc, ["Case", "Subject", "Analyst", "Date"], cols=4)
    doc.add_heading("Summary", level=1)
    note(doc, "One paragraph: the question and the answer.")
    doc.add_heading("Key findings", level=1)
    note(doc, CONFIDENCE)
    grid(doc, ["#", "Finding", "Confidence"], rows=5)
    doc.add_heading("Recommended action", level=1)
    note(doc, "What the reader should do now.")
    doc.add_heading("Subject at a glance", level=1)
    fields(doc, ["Name(s)", "Usernames", "Email addresses", "Phone numbers",
                 "Location", "Main profiles"])
    sources_annex(doc)
    save(doc, "Argos_Report_Executive.docx")


def subject_profile():
    doc = new_document("Subject Profile", "Identifiers and online footprint of one person or organisation")
    fields(doc, ["Case", "Profile of", "Analyst", "Last updated"], cols=4)
    doc.add_heading("Identity", level=1)
    grid(doc, ["Attribute", "Value", "Source (Ref)", "Confidence"], rows=8)
    note(doc, "Name, aliases, date of birth, gender, nationality, physical "
              "description, photographs (keep the image file in the case folder).")
    doc.add_heading("Addresses", level=1)
    grid(doc, ["Address", "Type (home, work, mail)", "Period", "Ref"], rows=3)
    doc.add_heading("Phone numbers", level=1)
    grid(doc, ["Number", "Carrier / type", "Period", "Ref"], rows=3)
    doc.add_heading("Email addresses", level=1)
    grid(doc, ["Address", "Provider", "Seen in breaches", "Ref"], rows=3)
    doc.add_heading("Online accounts", level=1)
    grid(doc, ["Platform", "Handle", "Profile URL", "Account ID", "Created / active", "Ref"], rows=8)
    doc.add_heading("Associates", level=1)
    grid(doc, ["Name", "Relationship", "How established", "Ref"], rows=4)
    doc.add_heading("Employment, education, organisations", level=1)
    grid(doc, ["Organisation", "Role", "Period", "Ref"], rows=3)
    doc.add_heading("Vehicles and assets", level=1)
    grid(doc, ["Item", "Details", "Ref"], rows=3)
    sources_annex(doc)
    save(doc, "Argos_Subject_Profile.docx")


def event_assessment():
    doc = new_document("Event Assessment", "Open-source monitoring of a planned or ongoing event")
    fields(doc, ["Event", "Assessment reference", "Date(s)", "Location", "Organiser",
                 "Expected attendance", "Analyst", "Status (draft / issued)"], cols=4)
    doc.add_heading("Assessment", level=1)
    note(doc, "Bottom line first: overall risk level and the main reasons. " + CONFIDENCE)
    doc.add_heading("Contacts", level=1)
    grid(doc, ["Role", "Name", "Phone", "Email / radio"], rows=4)
    doc.add_heading("What we monitor", level=1)
    grid(doc, ["Type (hashtag, account, page, group, channel)", "Value", "Platform", "Why"], rows=6)
    doc.add_heading("Groups and people of interest", level=1)
    grid(doc, ["Name", "Description", "Stance / intent", "Ref"], rows=4)
    doc.add_heading("Indicators and incidents", level=1)
    grid(doc, ["Time (UTC)", "What was seen", "Where", "Assessed significance", "Ref"], rows=6)
    doc.add_heading("Resources", level=1)
    note(doc, "Live feeds, maps, dashboards, browser tab groups used for the event.")
    sources_annex(doc)
    save(doc, "Argos_Event_Assessment.docx")


def threat_notice():
    doc = new_document("Threat Notice", "Short notice for chat, email or radio relay")
    fields(doc, ["Subject", "Time issued (UTC)", "Issued by", "Distribution"])
    doc.add_heading("Message", level=1)
    note(doc, "Write it so it can be pasted as plain text: WHAT was found, WHERE, "
              "WHEN, WHO is involved, HOW credible it is, WHAT you recommend. "
              "Keep it under 150 words.")
    fields(doc, ["What", "Where", "When", "Who", "Credibility", "Recommendation"])
    save(doc, "Argos_Threat_Notice.docx")


def cover_sheet():
    doc = new_document("Case Cover Sheet", "Opening record and chain of custody")
    fields(doc, ["Case", "Opened (UTC)", "Requested by", "Lead analyst",
                 "Authorisation / legal basis", "Handling / classification",
                 "Argos case folder", "Closed (UTC)"], cols=4)
    doc.add_heading("Tasking", level=1)
    note(doc, "The exact question, in the words of the requester.")
    doc.add_heading("Restrictions", level=1)
    checkboxes(doc, ["Passive collection only (no contact with the subject)",
                     "No research accounts may interact with the subject",
                     "No paid data sources", "Results must not leave this workstation",
                     "Other: ____________________"])
    doc.add_heading("Research identity and environment", level=1)
    fields(doc, ["Virtual machine / snapshot", "Network (VPN, Tor, other)",
                 "Research accounts used", "Argos version"])
    doc.add_heading("Chain of custody", level=1)
    note(doc, "Every hand-over of evidence (export, copy to media, delivery). "
              "Quote the SHA-256 from the run folder's SHA256SUMS.")
    grid(doc, ["Date / time (UTC)", "Item", "SHA-256", "From", "To", "Purpose"], rows=5)
    doc.add_heading("Closure", level=1)
    checkboxes(doc, ["Report issued", "Evidence archived with hashes verified",
                     "Research accounts reviewed", "VM snapshot reverted or archived"])
    save(doc, "Argos_Case_Cover_Sheet.docx")


def policy():
    doc = new_document("Online Investigations Policy", "Template for teams that research open sources")
    note(doc, "Adapt every section to your organisation and to the law that "
              "applies to you. This template is not legal advice.")
    sections = [
        ("1. Purpose", "Why the organisation carries out online research and what this policy protects: the people researched, the investigators and the integrity of the evidence."),
        ("2. Scope", "Who the policy applies to and which activities it covers (open sources, social media, commercial data, research accounts)."),
        ("3. Definitions", "Open source, publicly available information, private information, research account, run folder, evidence."),
        ("4. Lawful basis and proportionality", "When research is allowed, who authorises it, how necessity and proportionality are assessed and recorded."),
        ("5. Authorisation", "Levels of research (passive viewing, use of research accounts, interaction) and who may approve each."),
        ("6. Research accounts and identities", "How accounts are requested, created, recorded, used and retired; what they may and may not do."),
        ("7. Operational security", "Dedicated virtual machines, network separation, no personal accounts, handling of malicious content."),
        ("8. Collection and evidence", "Capture standards (time in UTC, URL, hash), storage in case folders, chain of custody."),
        ("9. Data protection and retention", "Minimisation, access control, retention periods, deletion, subject rights."),
        ("10. Training and supervision", "Required training before research, refreshers, supervision and audit."),
        ("11. Review", "How often the policy is reviewed and who owns it."),
    ]
    for title, text in sections:
        doc.add_heading(title, level=1)
        note(doc, text)
    doc.add_heading("Revision history", level=1)
    grid(doc, ["Version", "Date", "Author", "Changes"], rows=3)
    save(doc, "Argos_OSINT_Policy.docx")


if __name__ == "__main__":
    full_report()
    executive_report()
    subject_profile()
    event_assessment()
    threat_notice()
    cover_sheet()
    policy()
